"""Training & onboarding content: PowerPoint decks + orientation quiz.

Decks built here with python-pptx:
- ``worker`` deck: how to use the phone app and website, expected behavior in hosts'
  homes, ethics, sexual harassment policy, and other new-hire orientation items.
- ``host`` deck: how Broom Service works, services & pricing, the host portal, and
  funding/capital partners for leads.
- ``construction`` deck: Buildstack Construction crew orientation — OSHA-10 baseline
  (fatal four), PPE, ladder/scaffold, hazard communication, lockout/tagout, silica/dust/
  asbestos awareness, pay app & direct deposit, harassment policy, and a trades knowledge
  test.

Each deck is returned as in-memory ``.pptx`` bytes for storage in the document
library. The Broom worker quiz (10 questions) and the construction OSHA-10 / trades
quiz (``OSHA_QUIZ``) are defined here too.
"""

QUIZ = [
    {
        "q": "How do you get the worker app on your phone?",
        "options": [
            "Download it from the App Store",
            "Open the site on your phone and save it to your home screen",
            "It installs itself",
        ],
        "answer": 1,
    },
    {
        "q": "Why do you clock in and clock out on every job?",
        "options": [
            "So the office knows where you are",
            "Because that's how you get paid",
            "Just for fun",
        ],
        "answer": 1,
    },
    {
        "q": "When you arrive at a property, what should you do first?",
        "options": [
            "Clock in when you're inside the work area and then start cleaning",
            "Text your friends",
            "Wait outside until someone tells you to start",
        ],
        "answer": 0,
    },
    {
        "q": "Where do you find directions to your job?",
        "options": [
            "Call the office every time",
            "The map in the worker app",
            "Guess the address",
        ],
        "answer": 1,
    },
    {
        "q": "Can you eat food, use the bed, or take supplies from the host's home?",
        "options": [
            "Yes, if you're careful",
            "Only snacks",
            "No — it's the host's home; you treat everything with respect",
        ],
        "answer": 2,
    },
    {
        "q": "What should you do if you see damage or something broken at a property?",
        "options": [
            "Hide it",
            "Report it right away with a photo",
            "Fix it with tape",
        ],
        "answer": 1,
    },
    {
        "q": "How should you treat every person you meet at a property or office?",
        "options": [
            "With respect, always",
            "Only if they talk to you first",
            "Whatever, it doesn't matter",
        ],
        "answer": 0,
    },
    {
        "q": "Is unwelcome touching or sexual comments ever okay at work?",
        "options": [
            "No, never — it's harassment and won't be tolerated",
            "Yes if someone jokes about it",
            "Sometimes",
        ],
        "answer": 0,
    },
    {
        "q": "Where can you report a problem or concern safely?",
        "options": [
            "Nowhere",
            "Tell the owner or use the reporting channel — it's confidential",
            "Post it on social media",
        ],
        "answer": 1,
    },
    {
        "q": "Where do you find your paycheck and pay stubs?",
        "options": [
            "Under 'Pay' in the worker app or portal",
            "The host sends it by mail",
            "You have to ask in person",
        ],
        "answer": 0,
    },
]


OSHA_QUIZ = [
    {
        "q": "The 'Fatal Four' on construction sites are falls, struck-by, caught-in/between, and ______.",
        "options": ["Electrical", "Sun exposure", "Fatigue", "Loud noise"],
        "answer": 0,
        "topic": "OSHA-10 fatal four",
    },
    {
        "q": "At what height does fall protection (guardrails, harness + anchor) become required on most residential/commercial construction?",
        "options": [
            "3 feet",
            "6 feet",
            "15 feet",
            "Only on roofs",
        ],
        "answer": 1,
        "topic": "Fall protection",
    },
    {
        "q": "When working near or over moving equipment (forklift, excavator), what's the key struck-by protection?",
        "options": [
            "Stay in the equipment operator's blind spot",
            "Stay out of the swing/reach path and wear hi-vis so the operator sees you",
            "Wave your arms so they know to stop",
            "There's no way to prevent struck-by injuries",
        ],
        "answer": 1,
        "topic": "Struck-by",
    },
    {
        "q": "Which PPE is required on every active job site?",
        "options": [
            "Hard hat + eye protection",
            "Long sleeves only",
            "Steel-toe boots, nothing else",
            "PPE is optional for experienced crew",
        ],
        "answer": 0,
        "topic": "PPE",
    },
    {
        "q": "On a ladder, what is the correct body position?",
        "options": [
            "Lean backward while reaching",
            "Keep 3 points of contact and stay below the top rung",
            "Stand on the very top to reach high",
            "Face away from the ladder",
        ],
        "answer": 1,
        "topic": "Ladder safety",
    },
    {
        "q": "What does Lockout/Tagout (LOTO) prevent?",
        "options": [
            "Theft of tools",
            "Accidental energy release when working on electrical, plumbing, or equipment",
            "Getting paint on uniforms",
            "Nothing — it's just paperwork",
        ],
        "answer": 1,
        "topic": "Lockout/tagout",
    },
    {
        "q": "On a tape measure, the smallest marked increment is usually:",
        "options": ["1 inch", "1/2 inch", "1/16 inch", "2 inches"],
        "answer": 2,
        "topic": "Reading a tape measure",
    },
    {
        "q": "A 12-ft wall at 8 ft tall needs how many square feet of drywall?",
        "options": ["64 sq ft", "96 sq ft", "80 sq ft", "120 sq ft"],
        "answer": 1,
        "topic": "Simple trade math",
    },
    {
        "q": "If you work 46 hours in one week at a regular hourly rate, which hours are overtime?",
        "options": [
            "None",
            "Hours 41-46 (6 hours at 1.5x)",
            "The first 40, then rebuild the schedule",
            "Only Saturday hours",
        ],
        "answer": 1,
        "topic": "Payroll & Overtime",
    },
    {
        "q": "To control silica dust from cutting concrete or masonry, the best practice is:",
        "options": [
            "Wet cutting + dust collection + a respirator when required",
            "Just hold your breath",
            "Wear a bandana",
            "Cut faster so dust spreads less",
        ],
        "answer": 0,
        "topic": "Silica / dust awareness",
    },
    {
        "q": "Unwelcome touching, comments, or sexual jokes on the job site are:",
        "options": [
            "Fine if everyone laughs",
            "Harassment — zero tolerance, report confidentially, no retaliation",
            "Only a problem in the office",
            "Allowed between friends",
        ],
        "answer": 1,
        "topic": "Workplace harassment policy",
    },
    {
        "q": "Before working near electrical panels or exposed wires you should:",
        "options": [
            "Assume it's dead and start cutting",
            "LOTO it, verify with a tester, and use GFCI where needed",
            "Just avoid touching metal",
            "Stand on a metal ladder",
        ],
        "answer": 1,
        "topic": "Electrical safety",
    },
]


BRAND = {
    "title_color": (49, 46, 129),  # indigo-900
    "accent_color": (99, 102, 241),  # indigo-500
    "dark": (15, 23, 42),  # slate-900
    "light": (248, 250, 252),  # slate-50
}


def _add_title_slide(prs, doc_title, subtitle):
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    title = slide.shapes.title
    title.text = doc_title
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
    title.text_frame.paragraphs[0].runs[0].font.bold = True
    title.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(*BRAND["title_color"])
    if slide.shapes.placeholders and len(slide.shapes.placeholders) > 1:
        sub = slide.placeholders[1]
        sub.text = subtitle
        for para in sub.text_frame.paragraphs:
            for run in para.runs:
                run.font.size = Pt(18)
                run.font.color.rgb = RGBColor(*BRAND["dark"])
    return slide


def _add_bullets_slide(prs, doc_title, bullets):
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = doc_title
    title.text_frame.paragraphs[0].runs[0].font.bold = True
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    title.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(*BRAND["title_color"])
    body = slide.placeholders[1].text_frame
    body.clear()
    first = True
    for item in bullets:
        if first:
            para = body.paragraphs[0]
            first = False
        else:
            para = body.add_paragraph()
        para.text = item
        para.level = 0
        para.font.size = Pt(18)
        for run in para.runs:
            run.font.color.rgb = RGBColor(*BRAND["dark"])
    return slide


def _slides_content(kind: str) -> list[tuple[str, list[str]]]:
    if kind == "worker":
        return [
            ("Welcome to Broom Service!", [
                "Thank you for joining our cleaning & co-hosting crew",
                "Today: your app, your job, your pay, and how we act",
                "Short quiz at the end — easy if you pay attention",
            ]),
            ("About Our Company", [
                "We clean & co-host short-term rentals (Airbnb, Vrbo, direct)",
                "Guests fund the cleaning fee — hosts pay nothing out of pocket",
                "We win by doing 5-star turnovers, every time",
            ]),
            ("Your Phone App — Download & Login", [
                "Open https://bizstackperks.com on your phone",
                "Save it to your home screen so it opens like an app",
                "Log in with your email and the PIN the owner set for you",
                "That's it — no store download needed",
            ]),
            ("Your Phone App — Jobs & Map", [
                "Your assigned jobs show with the date, time and address",
                "Tap 'Directions' / the map to navigate to the property",
                "Jobs stay on your screen so you always know your day",
            ]),
            ("Clock In / Clock Out", [
                "When you arrive, the app checks you're at the right property",
                "Clock in when you're inside the work area and start",
                "Clock out when the job is finished (before you leave!)",
                "Clock in AND out on every job — that's what pays you",
            ]),
            ("Paychecks & Pay", [
                "You get paid per job at your rate on your paycheck",
                "Find pay stubs under 'Pay' in the app or on the web portal",
                "Questions about pay? Ask the office—never negotiate in the host's home",
            ]),
            ("Using the Website Portal", [
                "The website (https://bizstackperks.com) shows the same schedule",
                "Log in at /worker-login or through /login with your email + PIN",
                "See your jobs, map, and paychecks in one place",
            ]),
            ("House Rules in Hosts' Homes", [
                "You are a guest in their property — act like it",
                "No smoking, no eating host food, no using beds or bathrooms for personal use",
                "Leave every place exactly as you found it, plus cleaner",
                "Treat the host's stuff the way you'd want yours treated",
            ]),
            ("The 5-Star Turnover Checklist", [
                "Trash out, then all linen exchanged",
                "Bathrooms + kitchen sanitized top to bottom",
                "All floors vacuumed and mopped",
                "Dust everything, restock supplies, stage the space",
                "Photo-verify each room, then mark the job done",
            ]),
            ("Staying Safe at Work", [
                "Use supplies safely; report anything unsafe",
                "Let the office know any medical issue you have before jobs",
                "If you feel unsafe at a property, leave and call the office",
            ]),
            ("Workplace Ethics & Integrity", [
                "Be honest about time, vetting, and what you complete",
                "Don't cut corners to rush — 5-star is the brand",
                "Never share host or guest information with anyone",
            ]),
            ("Sexual Harassment Policy", [
                "Zero tolerance — always, toward anyone, at work",
                "Unwelcome touching, comments, jokes, or advances are prohibited",
                "Reporting is safe and confidential — no retaliation, ever",
            ]),
            ("Reporting Problems & Concerns", [
                "Damage, maintenance, host issues → report right away with a photo",
                "Harassment or anything wrong → tell the owner or use the report channel",
                "Reporting is never punished. Ever.",
            ]),
        ]
    if kind in ("construction", "construction-osha", "osha"):
        return [
            ("Welcome to Buildstack Construction!", [
                "Thank you for joining our licensed general contracting crew",
                "Today: your pay app, job site safety (OSHA-10 baseline), and how we act",
                "Short knowledge test at the end — pass to move to jobs",
            ]),
            ("About Our Company", [
                "Licensed, insured general contractor — resi + commercial",
                "Whole-home renovations, kitchens, baths, drywall, roofing, decks, all trades",
                "We build trust with 5-star work, safety first, every day",
            ]),
            ("Your Pay App — Login & Direct Deposit", [
                "Crew portal at your site login; use your phone + PIN",
                "Punch in/out on every job — that's what pays you",
                "Set up direct deposit (Stripe) so checks land automatically",
                "See hours, timesheets, and pay history under 'Pay'",
            ]),
            ("Clock In / Clock Out", [
                "Clock in when you're on site and ready to work",
                "Clock out when the job is done (before you leave)",
                "Clock in AND out on every job — that's how payroll happens",
                "Report all hours honestly; never punch for someone else",
            ]),
            ("OSHA-10: The Fatal Four", [
                "Falls — the #1 killer on job sites",
                "Struck-by — stay out of equipment swing/reach paths",
                "Caught-in/between — never enter trenches or between moving parts",
                "Electrical — LOTO, test before touching, use GFCI",
            ]),
            ("Fall Protection", [
                "Fall protection required at 6 feet and above",
                "Guardrails, covers, or a harness + anchor — no exceptions",
                "Never work on a roof without tie-off or other approved fall protection",
                "Keep ladders on firm ground: 3 points of contact, don't stand above the top rung",
            ]),
            ("PPE — Every Job, No Exceptions", [
                "Hard hat + eye protection on active job sites",
                "Gloves, steel-toe footwear, hearing protection, hi-vis as the task requires",
                "Check your PPE daily; report damaged gear to the owner",
                "No PPE, no work — period",
            ]),
            ("Hazard Communication & Chemicals", [
                "Read SDS before using any chemical product",
                "Eye/face wash and ventilation when mixing or spraying",
                "Never mix bleach with ammonia or other cleaners (deadly gas)",
                "Use the right chemical for the right task; store safely",
            ]),
            ("Lockout / Tagout (LOTO)", [
                "LOTO prevents accidental energy release — electrical, plumbing, equipment",
                "Only the person who applied the lock removes it",
                "Verify with a tester that circuits are dead before working",
                "Never wedge a switch or bypass a safety guard",
            ]),
            ("Trench & Excavation Safety", [
                "Never enter an unprotected trench — that's caught-in/between risk",
                "Trenches 5 ft+ need shoring, sloping, or a protective system",
                "Keep spoil piles and equipment clear of the trench edge",
                "When in doubt, ask the owner before going in",
            ]),
            ("Silica, Dust & Asbestos Awareness", [
                "Wet cut concrete/masonry + use dust collection",
                "Wear a respirator where silica or dust is heavy",
                "Older builds (pre-1970s) may have asbestos or lead paint — don't disturb it",
                "Report any suspected asbestos/lead to the owner immediately",
            ]),
            ("Workplace Ethics & HR", [
                "Be honest about time, materials, and what you complete",
                "Sexual harassment is zero tolerance — unwanted comments/touching, any time",
                "Report problems confidentially — no retaliation, ever",
                "Never share client or business information with anyone",
            ]),
            ("Your Knowledge Test", [
                "Short test: tape-measure reading, simple math, OSHA basics",
                "Covers framing, drywall, roofing, tile, plumbing basics",
                "Pass it and you're clear to move to jobs",
                "Certificate is recorded in your crew record",
            ]),
        ]
    if kind in ("construction-trades", "trades", "construction-tools"):
        return [
            ("Welcome to Tools of the Trade!", [
                "Today: the trades knowledge you'll use on real jobs",
                "Tape measure, framing, drywall, roofing, plumbing, electrical, HVAC, tile, painting, concrete",
                "Short test at the end — pass it and you're cleared for jobs",
            ]),
            ("About Buildstack Construction", [
                "Licensed, insured general contractor — resi + commercial",
                "Whole-home renovations, kitchens, baths, drywall, roofing, decks, all trades",
                "We build trust with 5-star work, safety first, every day",
            ]),
            ("Reading a Tape Measure", [
                "Know your marks: 1/16, 1/8, 1/4, 1/2 and whole inches",
                "Measure twice, cut once — re-check before any cut",
                "Read from the hook end; account for the hook's thickness",
                "Practice on a board until it's second nature",
            ]),
            ("Basic Math on the Job", [
                "Simple addition & subtraction for cuts and layout",
                "Fractions: half of 3/4\" is 3/8\" — know your fractions",
                "Square and level every layout (3-4-5 or a square)",
                "Ask before guessing — math mistakes cost materials and time",
            ]),
            ("Framing Basics", [
                "Framing is the skeleton — walls, floors, and roof structure",
                "Standard stud spacing is 16\" or 24\" on center",
                "Keep studs plumb, plates straight, and corners square",
                "Safety: hands clear of nail gun paths, wear eye protection",
            ]),
            ("Drywall Basics", [
                "Hang sheets flat, seams centered over studs",
                "Screw spacing roughly every 12\" along studs",
                "Tape and float joints in stages — thin coats build smooth",
                "Never leave unfinished mud for the next shift",
            ]),
            ("Roofing & Siding Basics", [
                "Work from the bottom up so every layer overlaps water flow",
                "Nail patterns and fasteners matter — follow spec",
                "Never work a roof without tie-off/fall protection at 6 ft+",
                "Keep ladders firm: 3 points of contact, don't overreach",
            ]),
            ("Plumbing Basics", [
                "Know the main shut-off and how to kill water in seconds",
                "PEX, copper, PVC — each needs the right fittings and glue",
                "Every connection leaks until it's tested — pressure test before closing walls",
                "Call the owner for anything suspect; no guess-plumbing",
            ]),
            ("Electrical Basics", [
                "LOTO before touching anything — test that circuits are dead",
                "Wires are color-coded — hot, neutral, ground, never mixed",
                "Use GFCI where required; never overload a circuit",
                "Leave electrical beyond basics to the licensed electrician",
            ]),
            ("HVAC Basics", [
                "Gas is dangerous — smell/leak test, call it in immediately",
                "Keep vents, returns and units clear of debris",
                "Filters, coils, and lines need gentle handling",
                "Never bypass a safety switch or limit",
            ]),
            ("Tile & Flooring Basics", [
                "Prep the subfloor/backer — flat matters more than anything",
                "Use leveling/thinset spec for the tile size",
                "Spacers keep grout lines straight; wipe grout before it hardens",
                "Let adhesives, mortar, and sealers cure per label",
            ]),
            ("Painting & Finishing Basics", [
                "Prep is 80% of a good finish — clean, sand, prime",
                "Ventilate and wear a respirator for solvent/spray",
                "Keep drop cloths flat — no trip hazards",
                "Two thin coats beat one thick coat, every time",
            ]),
            ("Concrete & Masonry Basics", [
                "Wet-cut to control silica dust — respirator when it's heavy",
                "Fresh concrete burns skin — flush with water, wear gloves/boots",
                "Stay clear of ready-mix backing — spotters use agreed signals",
                "Trenches 5 ft+ need shoring/sloping — never enter unprotected",
            ]),
            ("Materials, Prep & Cleanup", [
                "Keep materials dry, stacked, and off wet ground",
                "Clean as you go — a tidy site is a safe site",
                "Protect finished work with drop cloths and floor protection",
                "End of day: tools secured, trash out, site locked",
            ]),
            ("Your Trades Test", [
                "Coverage: tape measure, math, framing, drywall, roofing, tile, plumbing, electrical basics",
                "70% to pass — study the slides, then take it",
                "Pass and the certificate is recorded in your crew record",
                "Questions? Ask the office before you test",
            ]),
        ]
    if kind in ("construction-safety", "safety", "osha"):
        return [
            ("Welcome to Job Site Safety!", [
                "Today: the OSHA-10 baseline that keeps everyone working safe",
                "The Fatal Four, fall protection, PPE, chemicals, LOTO, trenches, silica & asbestos",
                "Short test at the end — pass to be cleared for jobs",
            ]),
            ("OSHA-10: The Fatal Four", [
                "Falls — the #1 killer on job sites",
                "Struck-by — stay out of equipment swing/reach paths",
                "Caught-in/between — never enter trenches or between moving parts",
                "Electrical — LOTO, test before touching, use GFCI",
            ]),
            ("Fall Protection", [
                "Fall protection required at 6 feet and above",
                "Guardrails, covers, or a harness + anchor — no exceptions",
                "Never work on a roof without tie-off or other approved fall protection",
                "Keep ladders on firm ground: 3 points of contact, don't stand above the top rung",
            ]),
            ("PPE — Every Job, No Exceptions", [
                "Hard hat + eye protection on active job sites",
                "Gloves, steel-toe footwear, hearing protection, hi-vis as the task requires",
                "Check your PPE daily; report damaged gear to the owner",
                "No PPE, no work — period",
            ]),
            ("Hazard Communication & Chemicals", [
                "Read the SDS before using any chemical product",
                "Eye/face wash and ventilation when mixing or spraying",
                "Never mix bleach with ammonia or other cleaners (deadly gas)",
                "Use the right chemical for the right task; store safely",
            ]),
            ("Lockout / Tagout (LOTO)", [
                "LOTO prevents accidental energy release — electrical, plumbing, equipment",
                "Only the person who applied the lock removes it",
                "Verify with a tester that circuits are dead before working",
                "Never wedge a switch or bypass a safety guard",
            ]),
            ("Trench & Excavation Safety", [
                "Never enter an unprotected trench — that's caught-in/between risk",
                "Trenches 5 ft+ need shoring, sloping, or a protective system",
                "Keep spoil piles and equipment clear of the trench edge",
                "When in doubt, ask the owner before going in",
            ]),
            ("Silica, Dust & Asbestos Awareness", [
                "Wet cut concrete/masonry + use dust collection",
                "Wear a respirator where silica or dust is heavy",
                "Older builds (pre-1970s) may have asbestos or lead paint — don't disturb it",
                "Report any suspected asbestos/lead to the owner immediately",
            ]),
            ("Electrical Safety", [
                "Test before touching — assume it's live until verified dead",
                "Use GFCI-protected circuits and extension cords",
                "Inspect cords for damage; never daisy-chain improperly",
                "Only the licensed electrician does the electrical work",
            ]),
            ("Ergonomics & Lifting", [
                "Lift with your legs, not your back; get help for heavy loads",
                "Two-person lifts for sheet goods and long runs",
                "Use carts, dollies and lifting tools when available",
                "Stretch and stay hydrated in the heat",
            ]),
            ("Incident Reporting", [
                "Report every injury, near miss, or unsafe condition — immediately",
                "First aid kit locations are posted; know where yours is",
                "Fatal Four concerns override the schedule — stop and speak up",
                "No retaliation for safety reports — ever",
            ]),
            ("Your Safety Test", [
                "Coverage: Fatal Four, falls, PPE, chemicals, LOTO, trenches, silica/asbestos, electrical",
                "70% to pass — study the slides, then take it",
                "Pass and the certificate is recorded in your crew record",
                "Safety is the job before every job",
            ]),
        ]
    if kind in ("construction-app", "app", "crew-app"):
        return [
            ("Welcome to the Crew App!", [
                "Today: how to get the app, log in, and track your time",
                "Download & install, log in, jobs & map, clock in/out, time reporting, direct deposit",
                "Short quiz at the end — easy if you follow along",
            ]),
            ("Download & Install the App", [
                "Open construction.bizstackperks.com on your phone",
                "Log in there the first time — it installs itself as a web app",
                "Use Chrome or Safari and choose 'Add to Home Screen'",
                "That's it — no app store download needed",
            ]),
            ("Add to Home Screen", [
                "In the browser menu, tap 'Add to Home Screen'",
                "It now opens full-screen, just like a native app",
                "Name it whatever you like — 'Buildstack' works",
                "The icon lives on your home screen for every shift",
            ]),
            ("Log In with Your Phone + PIN", [
                "Your login is your phone number and the PIN the owner set",
                "PINs are personal — never share yours, never log in for someone else",
                "Lock your phone when you're not using it",
                "Forgot your PIN? Call the office and they'll reset it",
            ]),
            ("Your Job List & Map", [
                "Your assigned jobs show with date, time, and address",
                "Tap 'Directions' / the map to navigate to the job",
                "Jobs stay on your screen so you always know your day",
                "See who's on the crew for that job",
            ]),
            ("Clock In / Clock Out", [
                "Clock in when you're on site and ready to work",
                "Clock out when the job is done — before you leave!",
                "Clock in AND out on every job — that's how payroll happens",
                "The app checks you're at the right property",
            ]),
            ("Time Reporting & Overtime", [
                "Report ALL hours honestly — every minute you worked",
                "Never punch in/out for someone else — that's a firing offense",
                "Overtime rules apply when they apply — the app tracks it",
                "Check your timesheet at end of week before payday",
            ]),
            ("Pay & Direct Deposit", [
                "Set up direct deposit (Stripe) so checks land automatically",
                "You get paid per job at your rate on your paycheck",
                "Pay stubs, hours, and history live under 'Pay'",
                "Pay questions go to the office — never the client",
            ]),
            ("Getting Help", [
                "App stuck? Restart it, then call the office",
                "Login issues → call the office, they'll reset your PIN",
                "Job questions, changes, or concerns → call the office or use the app",
                "The office is your backstop — 24/7 if needed",
            ]),
            ("Your App Quiz", [
                "Coverage: install, login, jobs/map, clock in/out, time reporting, pay",
                "70% to pass — follow the slides and you're set",
                "Pass and it's recorded in your crew record",
                "Welcome aboard — see you on the job",
            ]),
        ]
    if kind in ("construction-ethics", "ethics", "hr"):
        return [
            ("Welcome to Workplace Ethics!", [
                "Today: how we act, on and off the job site",
                "Honesty, respect, confidentiality, and zero tolerance for harassment",
                "Short test at the end — pass to be cleared for jobs",
            ]),
            ("Honesty Is the Foundation", [
                "Be honest about time, materials, and what you complete",
                "Never punch for someone else — report all hours truthfully",
                "Don't cut corners to rush — 5-star quality is the brand",
                "If you made a mistake, say so — fixing it early is what pros do",
            ]),
            ("Client Property & Trust", [
                "You work inside people's homes and businesses — act like a guest",
                "No smoking, no eating client food, no using their things",
                "Leave every site exactly as you found it, plus better",
                "Treat their property like you'd want yours treated",
            ]),
            ("Confidentiality", [
                "Never share client or business information with anyone",
                "No photos of client homes, interiors, or materials on social media",
                "Addresses, keys, codes, and pricing stay in-house",
                "Privacy is trust — lose it once, you lose everything",
            ]),
            ("Sexual Harassment — Zero Tolerance", [
                "Zero tolerance — always, toward anyone, at work",
                "Unwelcome touching, comments, jokes, or advances are prohibited",
                "Reporting is safe and confidential — no retaliation, ever",
                "See it? Report it. The policy protects everyone",
            ]),
            ("Reporting Problems & Concerns", [
                "Damage, hazards, or client issues → report right away with a photo",
                "Harassment or anything wrong → tell the owner or use the report channel",
                "Reporting is never punished. Ever.",
                "Silence is not loyalty — reporting protects the crew",
            ]),
            ("Drugs, Alcohol & Policy", [
                "No alcohol, drugs, or impairment on any job site — ever",
                "That includes the night before — show up ready to work safely",
                "Prescription meds that affect work must be disclosed to the office",
                "Safety policy violations are grounds for removal from jobs",
            ]),
            ("Company Tools & Materials", [
                "Company tools and materials are for jobs, not personal use",
                "Log and return equipment at end of day",
                "Report damaged or lost gear — don't hide it",
                "Stealing is termination and legal action — no exceptions",
            ]),
            ("Your Ethics Test", [
                "Coverage: honesty, time reporting, client trust, confidentiality, harassment, reporting",
                "70% to pass — study the slides, then take it",
                "Pass and the certificate is recorded in your crew record",
                "Ethics isn't a policy — it's who we are",
            ]),
        ]
    return [
        ("Welcome to Broom Service", [
            "Short-term rental cleaning + co-hosting, powered by automation",
            "You're the kind of host we love — a property with potential",
            "This deck shows services, pricing, your portal, and growth options",
        ]),
        ("What We Do", [
            "Turnover cleaning, deep cleaning, linen restock, inspections",
            "Digital co-hosting: 24/7 AI assistant, dynamic pricing, review help",
            "Full-service management: cleaning, vendors, multi-channel bookings",
            "All backed by photo-verified quality you can see",
        ]),
        ("Service Pricing", [
            "Turnover Cleaning $120  |  Deep Cleaning $200",
            "Linen Restock $50  |  Inspection $75",
            "Cleaning is GUEST-FUNDED — you pay $0 out of pocket",
            "Co-hosting: 10-15% (digital) or 20-30% (full-service) of gross bookings",
        ]),
        ("How Booking Works", [
            "Guests contact the 24/7 assistant by text or call",
            "Assistant books the slot and sends a secure Stripe payment link",
            "Guest pays at booking — your calendar and ledger update automatically",
        ]),
        ("Your Free Rental Analysis", [
            "Get a real, data-backed earnings report for your property",
            "Home value, market rent, nightly STR estimate, self-managed vs co-hosted",
            "Free via the site — just submit your address",
        ]),
        ("Your Host Portal", [
            "Log in at /host-login to see YOUR properties, bookings and payments",
            "A live map of your portfolio",
            "Only your data — ever",
        ]),
        ("Calendar & Payments", [
            "Bookings sync to your calendar",
            "Paid bookings appear as revenue automatically",
            "Unpaid guests can be re-sent a fresh payment link any time",
        ]),
        ("Growing With Financing", [
            "Need capital to furnish, stage, or convert your property?",
            "Broom Service connects eligible leads with bank/financing partners",
            "Funding is handled directly with the partner bank — we just connect you",
        ]),
        ("What Makes a 5-Star Host", [
            "Fast, friendly guest communication",
            "Immaculate turnover between guests (our job!)",
            "Clear house rules and a clean, stocked property",
        ]),
        ("Support & Next Steps", [
            "Questions? Text/call +1 (757) 846-9275 or email hello@bizstackperks.com",
            "Submit your property for a free analysis",
            "Let's get your listing earning",
        ]),
    ]


CONSTRUCTION_DECK_KINDS = {
    "construction": "construction",
    "construction-osha": "construction",
    "osha": "construction",
    "construction-trades": "construction-trades",
    "trades": "construction-trades",
    "construction-tools": "construction-trades",
    "construction-safety": "construction-safety",
    "safety": "construction-safety",
    "construction-app": "construction-app",
    "app": "construction-app",
    "crew-app": "construction-app",
    "construction-ethics": "construction-ethics",
    "ethics": "construction-ethics",
    "hr": "construction-ethics",
}


def deck_slides(kind: str) -> list[dict]:
    """Deck slides as {title, bullets, voice} — voice is the bot's narration text."""
    parsed = CONSTRUCTION_DECK_KINDS.get(kind, kind)
    if parsed not in ("construction", "construction-trades", "construction-safety", "construction-app", "construction-ethics"):
        parsed = "host" if parsed != "worker" else "worker"
    phrases = {
        "Welcome to Broom Service!": "Welcome to Broom Service.",
        "Welcome to Broom Service": "Welcome to Broom Service.",
        "About Our Company": "About our company.",
        "Welcome to Buildstack Construction!": "Welcome to Buildstack Construction.",
        "Welcome to Tools of the Trade!": "Welcome to Tools of the Trade.",
        "Welcome to Job Site Safety!": "Welcome to job site safety.",
        "Welcome to the Crew App!": "Welcome to the crew app.",
        "Welcome to Workplace Ethics!": "Welcome to workplace ethics.",
    }
    out = []
    for title, bullets in _slides_content(parsed):
        voice = phrases.get(title, f"Next up — {title}. ") + " "
        voice += " ".join(b for b in bullets if b)
        out.append({"title": title, "bullets": bullets, "voice": voice})
    return out


def build_deck(kind: str) -> bytes:
    """Build a powerpoint deck. kind: 'worker', 'host', or a construction deck
    ('construction'|'construction-trades'|'construction-safety'|'construction-app'|'construction-ethics').
    Returns .pptx bytes."""
    from pptx import Presentation

    kind = CONSTRUCTION_DECK_KINDS.get(kind, kind)
    if kind not in ("worker", "host", "construction",
                    "construction-trades", "construction-safety",
                    "construction-app", "construction-ethics"):
        kind = "host" if kind != "worker" else "worker"
    prs = Presentation()
    prs.slide_width = 12192000
    prs.slide_height = 6858000

    if kind == "worker":
        doc_title = "New Worker Orientation"
        subtitle = "Your app, your job, your pay, and how we do it — plus your quick quiz."
    elif kind == "construction":
        doc_title = "Crew Orientation + OSHA-10 Baseline"
        subtitle = "Your pay app, job-site safety (the fatal four), and how we build — plus your knowledge test."
    elif kind == "construction-trades":
        doc_title = "Tools of the Trade"
        subtitle = "Tape measure, framing, drywall, roofing, plumbing, electrical, HVAC, tile, painting, concrete — plus your trades test."
    elif kind == "construction-safety":
        doc_title = "Job Site Safety (OSHA-10 Baseline)"
        subtitle = "The Fatal Four, fall protection, PPE, chemicals, LOTO, trenches, silica & asbestos — plus your safety test."
    elif kind == "construction-app":
        doc_title = "Crew App — Install & Time Reporting"
        subtitle = "Download & install the app, log in, jobs & map, clock in/out, time reporting & overtime, direct deposit — plus your quiz."
    elif kind == "construction-ethics":
        doc_title = "Workplace Ethics & Professional Conduct"
        subtitle = "Honesty, client trust, confidentiality, zero tolerance for harassment, reporting — plus your ethics test."
    else:
        doc_title = "Host & Lead Onboarding"
        subtitle = "Welcome to Broom Service — how we help your property earn."

    _add_title_slide(prs, doc_title, subtitle)
    for title, bullets in _slides_content(kind):
        _add_bullets_slide(prs, title, bullets)

    return _save_pptx(prs)


def _save_pptx(prs) -> bytes:
    from io import BytesIO
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def grade_quiz(answers: list[int]) -> dict:
    """Grade quiz answers (list of 10 option indexes) -> score/percent + right/wrong."""
    if len(answers) != len(QUIZ):
        raise ValueError("expected 10 answers")
    correct = [i for i, a in enumerate(answers) if a == QUIZ[i]["answer"]]
    pct = round(100 * len(correct) / len(QUIZ))
    passed = pct >= 70
    return {
        "total": len(QUIZ),
        "correct": len(correct),
        "percent": pct,
        "passed": passed,
    }


def grade_osha_quiz(answers) -> dict:
    """Grade the OSHA-10 / trades orientation quiz.

    Accepts a list of {question_id, answer} dicts (answer = option index) or a flat
    list of option indexes. Returns pass/fail, percent, and missed topics.
    """
    if answers and isinstance(answers[0], dict):
        by_id = {}
        for item in answers:
            try:
                qid = int(item.get("question_id") or item.get("qid") or "0")
            except (TypeError, ValueError):
                qid = 0
            try:
                by_id[qid] = int(item.get("answer"))
            except (TypeError, ValueError):
                continue
        flat = [by_id.get(i) for i in range(len(OSHA_QUIZ))]
        answers = flat
    answers = list(answers or [])
    if len(answers) != len(OSHA_QUIZ):
        raise ValueError(f"expected {len(OSHA_QUIZ)} answers")
    correct = [i for i, a in enumerate(answers) if a == OSHA_QUIZ[i]["answer"]]
    missed = sorted({OSHA_QUIZ[i]["topic"] for i in range(len(OSHA_QUIZ)) if i not in correct})
    pct = round(100 * len(correct) / len(OSHA_QUIZ))
    passed = pct >= 70
    return {
        "total": len(OSHA_QUIZ),
        "correct": len(correct),
        "percent": pct,
        "passed": passed,
        "missed_topics": missed,
    }


# --- Trade skills tests -------------------------------------------------------
# Each test is informational company-sponsored training (self-check), not an
# official OSHA-10/30 certificate — that card can only be issued by an authorized
# OSHA trainer. Questions cover core trade knowledge + job-site safety rules.

TRADE_QUIZZES = {
    "electrician": {
        "title": "Electrician — Core Working Knowledge",
        "pass_percent": 70,
        "questions": [
            {
                "q": "Before working on a circuit, what must you do?",
                "options": [
                    "Assume it's dead and start cutting",
                    "Lock out / tag out, verify it's dead with a tester",
                    "Just wear rubber gloves",
                    "Turn off the lights",
                ],
                "answer": 1,
                "topic": "Electrical — LOTO & verification",
            },
            {
                "q": "Which breaker feeds the circuit when you're at a panel?",
                "options": [
                    "Whichever looks newest",
                    "Trace it with an appropriate tester or verify by labeling/checking",
                    "The top breaker, always",
                    "You don't need to know",
                ],
                "answer": 1,
                "topic": "Electrical — panel safety",
            },
            {
                "q": "A GFCI is required where?",
                "options": [
                    "Only outdoors",
                    "Damp/wet locations and where the code requires (kitchen, bath, outdoor)",
                    "Nowhere on new work",
                    "Only in basements",
                ],
                "answer": 1,
                "topic": "Electrical — GFCI",
            },
            {
                "q": "Loose wires in a junction box:",
                "options": [
                    "Are fine if taped tight",
                    "Must be terminated on proper connections and covered by a box cover",
                    "Can hang inside the wall",
                    "Only matter under load",
                ],
                "answer": 1,
                "topic": "Electrical — boxes & terminations",
            },
            {
                "q": "Which wire gauge is standard for a 15-amp branch circuit?",
                "options": ["10 AWG", "6 AWG", "14 AWG", "16 AWG"],
                "answer": 2,
                "topic": "Electrical — wire sizing",
            },
            {
                "q": "How do you confirm power is off before splicing?",
                "options": [
                    "Touch the wires quickly",
                    "Use a non-contact voltage tester AND/or a meter at the point of work",
                    "Check that lights are off",
                    "Ask a coworker to guess",
                ],
                "answer": 1,
                "topic": "Electrical — verify dead",
            },
            {
                "q": "A screwdriver with a metallic shaft used on live gear can:",
                "options": [
                    "Speed up the job",
                    "Short the circuit or shock you — use insulated tools and PPE",
                    "Never cause a problem",
                    "Only affect smart switches",
                ],
                "answer": 1,
                "topic": "Electrical — hand tools",
            },
            {
                "q": "After you finish a device install, the safest check is:",
                "options": [
                    "Turn power back on and test with proper equipment",
                    "Assume it works",
                    "Skip testing to save time",
                    "Have a homeowner test it first",
                ],
                "answer": 0,
                "topic": "Electrical — final test",
            },
        ],
    },
    "plumber": {
        "title": "Plumber — Core Working Knowledge",
        "pass_percent": 70,
        "questions": [
            {
                "q": "Before taking apart any fixture, you should:",
                "options": [
                    "Start wrenching immediately",
                    "Shut off the water supply and relieve pressure",
                    "Assume it's empty",
                    "Open every valve",
                ],
                "answer": 1,
                "topic": "Plumbing — isolate supply",
            },
            {
                "q": "A P-trap is primarily there to:",
                "options": [
                    "Collect gold",
                    "Seal against sewer gas and catch small debris",
                    "Slow down draining",
                    "Connect to the vent only",
                ],
                "answer": 1,
                "topic": "Plumbing — traps",
            },
            {
                "q": "When working on gas or pressurized lines, the first step is:",
                "options": [
                    "Cut the line",
                    "Lock out / shut off service and verify zero pressure",
                    "Smell check only",
                    "Nothing special",
                ],
                "answer": 1,
                "topic": "Plumbing — LOTO / gas",
            },
            {
                "q": "Drain lines need proper slope so they:",
                "options": [
                    "Look straight",
                    "Flow by gravity (typically 1/4 inch per foot) and self-clean",
                    "Sound quiet",
                    "Fill with water",
                ],
                "answer": 1,
                "topic": "Plumbing — drainage pitch",
            },
            {
                "q": "Flux + heat on copper joints:",
                "options": [
                    "Needs a fire watch and heat shield near combustibles",
                    "Is always safe",
                    "Produces no fumes",
                    "Should be done near drywall butts",
                ],
                "answer": 0,
                "topic": "Plumbing — brazing/soldering fire safety",
            },
            {
                "q": "A water heater's temperature & pressure (T&P) relief valve:",
                "options": [
                    "Can be capped off",
                    "Must be installed and discharge safely — never cap or block it",
                    "Is optional",
                    "Only goes on gas units",
                ],
                "answer": 1,
                "topic": "Plumbing — water heater safety",
            },
            {
                "q": "Before soldering, why must the line be dry?",
                "options": [
                    "Water conducts heat and prevents a good solder bond",
                    "It tastes better",
                    "To cool the torch",
                    "No reason",
                ],
                "answer": 0,
                "topic": "Plumbing — soldering prep",
            },
            {
                "q": "A water leak detected in a wall should be:",
                "options": [
                    "Covered with drywall patch",
                    "Reported, isolated (shut off supply), and dried/repaired properly",
                    "Ignored if small",
                    "Bleached",
                ],
                "answer": 1,
                "topic": "Plumbing — leak response",
            },
        ],
    },
    "roofing": {
        "title": "Roofing & Siding — Core Working Knowledge",
        "pass_percent": 70,
        "questions": [
            {
                "q": "Working on a roof without approved fall protection:",
                "options": [
                    "Is fine on low slopes",
                    "Is a violation at 6 feet+ — use a harness + anchor or guardrails/covers",
                    "Only matters in rain",
                    "Is allowed if you're experienced",
                ],
                "answer": 1,
                "topic": "Roofing — fall protection",
            },
            {
                "q": "On a ladder to the roof, you should:",
                "options": [
                    "Stand on the top rung",
                    "Keep 3 points of contact and extend it 3 feet above the roof edge",
                    "Lean sideways to reach",
                    "Carry tools in both hands",
                ],
                "answer": 1,
                "topic": "Roofing — ladder safety",
            },
            {
                "q": "When the forecast calls for high wind, roof work should be:",
                "options": [
                    "Continue normally",
                    "Stopped / deferred and materials secured",
                    "Done only by the fastest crew",
                    "Moved to the peak",
                ],
                "answer": 1,
                "topic": "Roofing — weather",
            },
            {
                "q": "A nail gun should always be:",
                "options": [
                    "Pointed at co-workers for fun",
                    "Treated as loaded — keep fingers clear, wear eye protection, never point at anyone",
                    "Dry-fired at feet",
                    "Carried by the trigger",
                ],
                "answer": 1,
                "topic": "Roofing — nail gun safety",
            },
            {
                "q": "Two workers both at the eave working below the other:",
                "options": [
                    "Can throw debris from below",
                    "Keep separated or use a gap safety plan so tools/debris don't hit anyone",
                    "Should race",
                    "Is standard practice",
                ],
                "answer": 1,
                "topic": "Roofing — struck-by / debris",
            },
            {
                "q": "Removed shingles loaded into a chute or dumpster must be:",
                "options": [
                    "Thrown over the roof edge",
                    "Kept in one controlled discharge point with a clear zone below",
                    "Left on the lawn",
                    "Buried onsite",
                ],
                "answer": 1,
                "topic": "Roofing — debris control",
            },
            {
                "q": "Underlayment should be installed so that:",
                "options": [
                    "Any order works",
                    "Upper laps shed water over lower (overlap with headlap as specified)",
                    "It drapes loose",
                    "It's only on the peak",
                ],
                "answer": 1,
                "topic": "Roofing — underlayment",
            },
            {
                "q": "Siding nails that are over-driven:",
                "options": [
                    "Are irrelevant",
                    "Can split panels and void warranty — set the nail flush/at spec",
                    "Make it stronger",
                    "Only affect color",
                ],
                "answer": 1,
                "topic": "Siding — fastening",
            },
        ],
    },
    "hvac": {
        "title": "HVAC — Core Working Knowledge",
        "pass_percent": 70,
        "questions": [
            {
                "q": "Before servicing an HVAC unit, you should:",
                "options": [
                    "Reach inside the cabinet",
                    "Lock out electrical power and verify it's disconnected",
                    "Assume it's inert",
                    "Only turn the thermostat off",
                ],
                "answer": 1,
                "topic": "HVAC — LOTO",
            },
            {
                "q": "Refrigerant must never be:",
                "options": [
                    "Recovered",
                    "Vented to atmosphere — recover it with proper equipment",
                    "Handled by a tech",
                    "Measured",
                ],
                "answer": 1,
                "topic": "HVAC — refrigerant handling",
            },
            {
                "q": "A capacitor on a blower/compressor can:",
                "options": [
                    "Hold a dangerous charge — discharge it with a proper resistor/bleeder, not your hand",
                    "Never shock anyone",
                    "Be touched safely",
                    "Only affect the fan",
                ],
                "answer": 0,
                "topic": "HVAC — capacitor discharge",
            },
            {
                "q": "Airflow restrictions (dirty filter, blocked coil) cause:",
                "options": [
                    "Nothing",
                    "Higher energy use, freezing coils, and compressor damage",
                    "Better airflow",
                    "Quieter operation",
                ],
                "answer": 1,
                "topic": "HVAC — airflow",
            },
            {
                "q": "Supply and return duct connections should be:",
                "options": [
                    "Sealed and insulated where required to prevent air loss and condensation",
                    "Left open to the attic",
                    "Taped with duct tape only",
                    "Optional",
                ],
                "answer": 0,
                "topic": "HVAC — ducts",
            },
            {
                "q": "Checking a flame or gas leak requires you to:",
                "options": [
                    "Smell only",
                    "Use proper detection/gauges and follow the unit's service procedure",
                    "Light a match",
                    "Ignore it",
                ],
                "answer": 1,
                "topic": "HVAC — gas safety",
            },
            {
                "q": "Condensate drain lines that are clear:",
                "options": [
                    "Prevent water damage and mold — keep them unobstructed",
                    "Are a problem",
                    "Only matter in winter",
                    "Should be capped",
                ],
                "answer": 0,
                "topic": "HVAC — condensate",
            },
            {
                "q": "When working on a roof-mounted unit:",
                "options": [
                    "No fall protection needed",
                    "Use fall protection and secure tools/scooters so nothing falls on people below",
                    "Sit on the condenser",
                    "Skip PPE",
                ],
                "answer": 1,
                "topic": "HVAC — rooftop fall protection",
            },
        ],
    },
    "painting": {
        "title": "Painting & Finishing — Core Working Knowledge",
        "pass_percent": 70,
        "questions": [
            {
                "q": "Homes built before 1978 may contain:",
                "options": [
                    "Nothing hazardous",
                    "Lead-based paint — don't sand/scrape without proper containment and PPE",
                    "Radioactive paint",
                    "Mercury only",
                ],
                "answer": 1,
                "topic": "Painting — lead awareness",
            },
            {
                "q": "Sanding drywall or primer creates:",
                "options": [
                    "No dust",
                    "Respirable dust — wear a respirator and ventilate/contain the area",
                    "Pure oxygen",
                    "A nice finish",
                ],
                "answer": 1,
                "topic": "Painting — dust control",
            },
            {
                "q": "Solvent-based coatings and spray applications need:",
                "options": [
                    "No ventilation",
                    "Ventilation, low-VOC where possible, and respirator protection",
                    "A heater",
                    "Only a hat",
                ],
                "answer": 1,
                "topic": "Painting — VOC & ventilation",
            },
            {
                "q": "On a ladder painting a high wall, you should:",
                "options": [
                    "Reach far sideways",
                    "Keep the ladder close, your body square, and move the ladder instead of overreaching",
                    "Stand on the top rung",
                    "Stretch the roller length instead",
                ],
                "answer": 1,
                "topic": "Painting — ladder safety",
            },
            {
                "q": "Surface must be free of dirt, dust, and loose paint because:",
                "options": [
                    "It looks better that day",
                    "Paint will not adhere well otherwise — prep is 80% of a good finish",
                    "It's faster",
                    "Of the smell",
                ],
                "answer": 1,
                "topic": "Painting — prep",
            },
            {
                "q": "Torn drop cloths under foot while painting can cause:",
                "options": [
                    "Nothing",
                    "Trip hazards — keep floors clear and cloths flat/taped",
                    "Better coverage",
                    "Static",
                ],
                "answer": 1,
                "topic": "Painting — slips/trips",
            },
            {
                "q": "Cleanup of latex paint:",
                "options": [
                    "Pour it down the drain",
                    "Wash brushes in water, collect rinsate, and dispose of paint properly — never down drains/storm",
                    "Bury the cans",
                    "Leave to dry in the sink",
                ],
                "answer": 1,
                "topic": "Painting — disposal",
            },
            {
                "q": "Two coats of primer on raw drywall:",
                "options": [
                    "Waits material",
                    "Seals joints and gives an even finish coat — apply per spec",
                    "Causes peeling",
                    "Is futile",
                ],
                "answer": 1,
                "topic": "Painting — finishing",
            },
        ],
    },
    "concrete": {
        "title": "Concrete & Masonry — Core Working Knowledge",
        "pass_percent": 70,
        "questions": [
            {
                "q": "Cutting concrete generates:",
                "options": [
                    "Invisible air",
                    "Silica dust — wet-cut and/or use dust collection plus a respirator",
                    "No hazard",
                    "Only noise",
                ],
                "answer": 1,
                "topic": "Concrete — silica",
            },
            {
                "q": "When a ready-mix truck is backing up:",
                "options": [
                    "Stand behind the truck to guide it",
                    "Stay clear — spotters use the agreed signals and no one stands in the swing path",
                    "Wave at the driver",
                    "It doesn't matter",
                ],
                "answer": 1,
                "topic": "Concrete — struck-by / backing",
            },
            {
                "q": "Freshly placed concrete in contact with skin:",
                "options": [
                    "Is harmless",
                    "Can cause chemical burns — flush with water immediately and wear gloves/boots",
                    "Tans you",
                    "Only stings on the legs",
                ],
                "answer": 1,
                "topic": "Concrete — chemical burns",
            },
            {
                "q": "Standing in a trench for utility work requires:",
                "options": [
                    "Nothing special",
                    "A protected system (shoring/sloping) per OSHA at 5 ft+ and a safe entry/exit",
                    "Sprinting in",
                    "Just a hard hat",
                ],
                "answer": 1,
                "topic": "Concrete — trench/excavation",
            },
            {
                "q": "Concrete cures best when:",
                "options": [
                    "Left to dry fast",
                    "Kept moist/cured per spec for several days to reach design strength",
                    "Baked",
                    "Frozen",
                ],
                "answer": 1,
                "topic": "Concrete — curing",
            },
            {
                "q": "A metal trowel floating before the mix is ready:",
                "options": [
                    "Speeds the pour",
                    "Can seal the surface and cause dusting/delamination — wait for the right set",
                    "Is required",
                    "Helps the color",
                ],
                "answer": 1,
                "topic": "Concrete — finishing",
            },
            {
                "q": "Formwork should be:",
                "options": [
                    "Reused until it fails",
                    "Braced/secured per the tie/brace plan — never stand under unsupported forms",
                    "Covered in oil on the job",
                    "Left loose",
                ],
                "answer": 1,
                "topic": "Concrete — forms & bracing",
            },
            {
                "q": "After the pour, foot traffic and rain on fresh concrete:",
                "options": [
                    "Are fine",
                    "Must be protected — keep off and cover until the mix is set",
                    "Help it dry",
                    "Only matters indoors",
                ],
                "answer": 1,
                "topic": "Concrete — fresh pour protection",
            },
        ],
    },
}

# Slug -> test definition used by the portal and phone app.
# Entry tests are short, simple, trade-specific comprehension checks (with
# free-text tape-measure reading questions where the trade reads a tape).
ENTRY_QUIZZES = {
    "tape": {
        "title": "Reading a Tape Measure — Entry",
        "pass_percent": 70,
        "questions": [
            {"q": "The small marks on a standard tape measure are _______ of an inch.",
             "options": ["1/8 or 1/16", "The same as inches", "Only decorative", "1 yard each"],
             "answer": 0, "topic": "Tape measure"},
            {"q": "How many inches are in one foot?", "options": ["10", "12", "16", "100"],
             "answer": 1, "topic": "Units"},
            {"q": "Write the measurement: the tape reads exactly 4 inches plus three of the 1/8-inch marks.",
             "type": "text", "accept": ["4 3/8", "4 3/8\"", "4-3/8", "4.375", "4 3/8 in", "4 3/8 inches"],
             "topic": "Tape measure"},
            {"q": "Write the measurement: the tape reads exactly 1 foot plus half of the next inch.",
             "type": "text", "accept": ["1' 1/2\"", "1 ft 1/2 in", "12 1/2", "12.5", "12 1/2\"", "1 foot 1/2 inch", "12 1/2 inches", "1' 0.5\"", "1 ft 0.5 in"],
             "topic": "Tape measure"},
            {"q": "How many inches is 2 feet?", "options": ["12", "20", "24", "36"], "answer": 2, "topic": "Units"},
            {"q": "Which is larger, 3/4 inch or 5/8 inch?", "options": ["3/4 inch", "5/8 inch", "They are equal", "Neither"],
             "answer": 0, "topic": "Fractions in the trades"},
        ],
    },
    "framing": {
        "title": "Framing — Entry",
        "pass_percent": 70,
        "questions": [
            {"q": "Standard wall stud spacing is:", "options": ["12 inches or 24 inches only", "16 or 24 inches on center", "Whatever looks right", "Every 10 inches"],
             "answer": 1, "topic": "Framing basics"},
            {"q": "A stud should be checked with a level so it is:", "options": ["Plumb", "Painted", "Rounded", "Loose"],
             "answer": 0, "topic": "Framing basics"},
            {"q": "Write the measurement: a 2x4 measured from end to end reads 8 feet exactly.", "type": "text",
             "accept": ["8 ft", "8'", "8 feet", "96 inches", "96 in", "8ft", "8 ft.", "8 feet long"],
             "topic": "Framing basics"},
            {"q": "When using a nail gun, your free hand should be:", "options": ["In the nail path", "Clear of the nail path and trigger", "On the trigger too", "Behind the board"],
             "answer": 1, "topic": "Framing safety"},
            {"q": "What is a common floor joist spacing?", "options": ["16 or 24 inches on center", "Always 7 feet", "No pattern", "Every 3 inches"],
             "answer": 0, "topic": "Framing basics"},
            {"q": "Framing walls should be checked so they are:", "options": ["Square and level", "Tilted slightly", "Random", "Not attached"],
             "answer": 0, "topic": "Framing basics"},
        ],
    },
    "drywall": {
        "title": "Drywall — Entry",
        "pass_percent": 70,
        "questions": [
            {"q": "Standard drywall for walls is usually:", "options": ["1/2 inch thick", "4 inches thick", "Paper thin", "1 inch thick only"],
             "answer": 0, "topic": "Drywall basics"},
            {"q": "Drywall screws should be spaced about every ______ along the studs.", "options": ["12 inches", "10 feet", "3 inches", "Only at corners"],
             "answer": 0, "topic": "Drywall basics"},
            {"q": "Seams between drywall sheets should:", "options": ["Land on a stud", "Float in mid-air", "Overlap randomly", "Touch the floor only"],
             "answer": 0, "topic": "Drywall basics"},
            {"q": "Write the measurement: a drywall sheet measures 48 inches wide.", "type": "text",
             "accept": ["48 in", "4 ft", "4 feet", "48\"", "48 inches", "4ft"],
             "topic": "Drywall basics"},
            {"q": "Tape and joint compound is applied:", "options": ["In thin coats, in stages", "All at once, thick", "Never", "Only on floors"],
             "answer": 0, "topic": "Drywall finishing"},
            {"q": "Which room usually gets moisture-resistant (green) drywall?", "options": ["Bathroom", "Living room only", "Garage only", "Attic"],
             "answer": 0, "topic": "Drywall moisture"},
        ],
    },
    "roofing": {
        "title": "Roofing & Siding — Entry",
        "pass_percent": 70,
        "questions": [
            {"q": "Fall protection is required when working at what height above a lower level?", "options": ["6 feet", "15 feet", "1 foot", "Only on 2-story roofs"],
             "answer": 0, "topic": "Roofing safety"},
            {"q": "Shingles are installed:", "options": ["From the bottom of the roof up", "From the peak down", "Any order", "Only on the ridge"],
             "answer": 0, "topic": "Roofing installation"},
            {"q": "Roof ladders and climbing should keep:", "options": ["3 points of contact", "2 points of contact", "No contact", "One hand full of tools"],
             "answer": 0, "topic": "Roofing safety"},
            {"q": "Write the measurement: a piece of siding is 12 feet long.", "type": "text",
             "accept": ["12 ft", "12'", "12 feet", "144 inches", "144 in", "12ft"],
             "topic": "Roofing & siding"},
            {"q": "Siding nails that are over-driven can:", "options": ["Split the panel", "Strengthen it", "Change the color", "Nothing"],
             "answer": 0, "topic": "Siding fastening"},
            {"q": "Underlayment (felt/ice-and-water) should be installed so it:", "options": ["Laps so water sheds downward", "Hangs loose", "Covers the peak only", "Wraps the gutters"],
             "answer": 0, "topic": "Roofing underlayment"},
        ],
    },
    "plumbing": {
        "title": "Plumbing — Entry",
        "pass_percent": 70,
        "questions": [
            {"q": "Before starting plumbing work, the first thing to know is:", "options": ["Where the shut-off is", "The color of the pipes", "How long the job is", "The apprentice's name"],
             "answer": 0, "topic": "Plumbing basics"},
            {"q": "PEX, copper, and PVC each need:", "options": ["The right fittings and tool for that material", "Are all the same", "No tools", "Only glue"],
             "answer": 0, "topic": "Plumbing materials"},
            {"q": "Before closing up walls, you should:", "options": ["Pressure-test the pipes", "Skip the test", "Pour water everywhere", "Paint the pipes"],
             "answer": 0, "topic": "Plumbing testing"},
            {"q": "Write the measurement: a length of pipe measures 3 feet 6 inches.", "type": "text",
             "accept": ["3 ft 6 in", "3'6\"", "3' 6\"", "42 inches", "42 in", "3 foot 6 inches", "3.5 ft", "3.5 feet"],
             "topic": "Plumbing"},
            {"q": "If you are unsure about a plumbing connection, you should:", "options": ["Ask the owner before guessing", "Guess and move on", "Hide it", "Blame the pipe"],
             "answer": 0, "topic": "Plumbing basics"},
            {"q": "What color pipe indicates a water SUPPLY line in most homes?", "options": ["Copper or PEX", "Always black gas pipe", "Always the drain", "Doesn't matter"],
             "answer": 0, "topic": "Plumbing basics"},
        ],
    },
    "electrical": {
        "title": "Electrical — Entry",
        "pass_percent": 70,
        "questions": [
            {"q": "Before working near wires or panels, you must:", "options": ["Assume it's live, then LOTO and test it's dead", "Start cutting", "Touch it quickly", "Pour water on it"],
             "answer": 0, "topic": "Electrical safety"},
            {"q": "GFCI outlets are required in:", "options": ["Damp/wet areas like kitchens & baths", "Bedrooms only", "Attics only", "Nowhere"],
             "answer": 0, "topic": "Electrical safety"},
            {"q": "If a cord looks frayed or damaged, you should:", "options": ["Stop using it and replace it", "Tape it loosely and continue", "Ignore it", "Wrap it in foil"],
             "answer": 0, "topic": "Electrical safety"},
            {"q": "Write the measurement: a 12-gauge wire run measures 50 feet.", "type": "text",
             "accept": ["50 ft", "50'", "50 feet", "600 inches", "50ft"],
             "topic": "Electrical"},
            {"q": "Real electrical work (panels, wiring) is done by:", "options": ["The licensed electrician", "Anyone with pliers", "Only if it's a weekend", "The homeowner, always"],
             "answer": 0, "topic": "Electrical scope"},
            {"q": "When testing a circuit to confirm it's dead, you use a:", "options": ["Voltage tester/meter", "Hammer", "Flashlight", "Wet finger"],
             "answer": 0, "topic": "Electrical testing"},
        ],
    },
    "hvac": {
        "title": "HVAC — Entry",
        "pass_percent": 70,
        "questions": [
            {"q": "Before servicing an HVAC unit, first:", "options": ["Lock out power and verify it's off", "Open the cabinet and look", "Assume it's off", "Spray it down"],
             "answer": 0, "topic": "HVAC safety"},
            {"q": "Refrigerant must never be:", "options": ["Vented to the air — it's recovered properly", "Measured", "Handled by techs", "Kept in the truck"],
             "answer": 0, "topic": "HVAC refrigerant"},
            {"q": "A blocked/dirty filter causes:", "options": ["Higher energy use and frozen coils", "Better airflow", "Nothing", "Quieter operation"],
             "answer": 0, "topic": "HVAC airflow"},
            {"q": "Write the measurement: a duct run measures 9 feet 4 inches.", "type": "text",
             "accept": ["9 ft 4 in", "9'4\"", "9' 4\"", "112 inches", "112 in", "9 foot 4 inches"],
             "topic": "HVAC"},
            {"q": "If you smell gas near HVAC equipment, you should:", "options": ["Stop, leave it closed, and call it in", "Keep working", "Light a match to check", "Ignore it"],
             "answer": 0, "topic": "HVAC gas safety"},
            {"q": "Never bypass a ______ on the unit.", "options": ["Safety switch or limit", "Filter", "Screw", "Label"],
             "answer": 0, "topic": "HVAC safety"},
        ],
    },
    "tile": {
        "title": "Tile & Flooring — Entry",
        "pass_percent": 70,
        "questions": [
            {"q": "The most important thing for tile is:", "options": ["A flat, prepped surface", "A bright color", "Extra glue", "No backer"],
             "answer": 0, "topic": "Tile prep"},
            {"q": "Tile spacers are used to:", "options": ["Keep grout lines straight and even", "Hold tools", "Level furniture", "Nothing"],
             "answer": 0, "topic": "Tile installation"},
            {"q": "In wet areas, tile goes on:", "options": ["Cement backer board, not regular drywall", "Paper", "Loose boards", "Carpet"],
             "answer": 0, "topic": "Tile backer"},
            {"q": "Write the measurement: a tile measures 12 inches by 24 inches.", "type": "text",
             "accept": ["12 x 24", "12 by 24", "12\" x 24\"", "1 ft x 2 ft", "1 foot by 2 feet", "12x24"],
             "topic": "Tile & flooring"},
            {"q": "After laying tile, you let the mortar/thinset:", "options": ["Cure before walking and grouting", "Stay wet", "Dry cracked", "Cover it in water"],
             "answer": 0, "topic": "Tile curing"},
            {"q": "Grout should be wiped from the tile surface:", "options": ["Before it hardens", "A week later", "Never", "With a hammer"],
             "answer": 0, "topic": "Tile grout"},
        ],
    },
    "painting": {
        "title": "Painting & Finishing — Entry",
        "pass_percent": 70,
        "questions": [
            {"q": "What is 80% of a good paint finish?", "options": ["Prep (clean, sand, prime)", "Buying expensive brushes", "Painting fast", "Lots of paint in one coat"],
             "answer": 0, "topic": "Painting prep"},
            {"q": "When spraying or using solvent paints, you need:", "options": ["Ventilation and a respirator", "A heater", "No protection", "Less time"],
             "answer": 0, "topic": "Painting VOC"},
            {"q": "Two thin coats are better than:", "options": ["One thick coat", "No coats", "Ten coats of primer", "Rolling once"],
             "answer": 0, "topic": "Painting finishing"},
            {"q": "Write the measurement: a wall to paint measures 10 feet by 8 feet.", "type": "text",
             "accept": ["10 x 8", "10 by 8", "10' x 8'", "10 ft x 8 ft", "10 x 8 ft", "10 feet by 8 feet"],
             "topic": "Painting"},
            {"q": "Drop cloths in the work area should be:", "options": ["Flat and taped (no trip hazards)", "Bunched up", "Wet", "Moved constantly"],
             "answer": 0, "topic": "Painting safety"},
            {"q": "Latex paint cleanup (brushes/rinsate) should:", "options": ["Never go down drains/storm — collect and dispose properly", "Go down the sink", "Be buried", "Stay on the brush"],
             "answer": 0, "topic": "Painting disposal"},
        ],
    },
    "concrete": {
        "title": "Concrete & Masonry — Entry",
        "pass_percent": 70,
        "questions": [
            {"q": "Cutting concrete creates:", "options": ["Silica dust — wet-cut & collect dust", "No risk", "Only noise", "Extra strength"],
             "answer": 0, "topic": "Concrete silica"},
            {"q": "Fresh concrete on skin can:", "options": ["Cause chemical burns — flush with water", "Help your hands", "Nothing", "Tan you"],
             "answer": 0, "topic": "Concrete safety"},
            {"q": "When a ready-mix truck backs up, you should:", "options": ["Stay clear; spotters use signals", "Stand behind it to guide", "Wave it in close", "Run next to it"],
             "answer": 0, "topic": "Concrete struck-by"},
            {"q": "Write the measurement: a concrete slab is 5 feet 6 inches wide.", "type": "text",
             "accept": ["5 ft 6 in", "5'6\"", "5' 6\"", "66 inches", "66 in", "5 foot 6 inches", "5.5 ft"],
             "topic": "Concrete"},
            {"q": "Standing in a trench for utility work 5 ft+ requires:", "options": ["Shoring or sloping protection", "Just a hard hat", "Sprinting in", "Nothing"],
             "answer": 0, "topic": "Concrete trench"},
            {"q": "Fresh concrete before it sets should be:", "options": ["Kept off and covered", "Walked on to test", "Sprayed with gasoline", "Left open to rain"],
             "answer": 0, "topic": "Concrete fresh pour"},
        ],
    },
}


# Slug -> test definition used by the portal and phone app.
ALL_TRAININGS = [
    {
        "slug": "osha",
        "title": "Site Safety Orientation (OSHA-10 baseline)",
        "pass_percent": 70,
        "questions": OSHA_QUIZ,
    },
    {"slug": "electrician", "title": "Electrician — Core Working Knowledge", "pass_percent": 70, "questions": TRADE_QUIZZES["electrician"]["questions"]},
    {"slug": "plumber", "title": "Plumber — Core Working Knowledge", "pass_percent": 70, "questions": TRADE_QUIZZES["plumber"]["questions"]},
    {"slug": "roofing", "title": "Roofing & Siding — Core Working Knowledge", "pass_percent": 70, "questions": TRADE_QUIZZES["roofing"]["questions"]},
    {"slug": "hvac", "title": "HVAC — Core Working Knowledge", "pass_percent": 70, "questions": TRADE_QUIZZES["hvac"]["questions"]},
    {"slug": "painting", "title": "Painting & Finishing — Core Working Knowledge", "pass_percent": 70, "questions": TRADE_QUIZZES["painting"]["questions"]},
    {"slug": "concrete", "title": "Concrete & Masonry — Core Working Knowledge", "pass_percent": 70, "questions": TRADE_QUIZZES["concrete"]["questions"]},
]

ALL_TRAININGS.extend(
    {"slug": "entry-" + slug, "title": t["title"], "pass_percent": t["pass_percent"], "questions": t["questions"]}
    for slug, t in ENTRY_QUIZZES.items()
)

TRAINING_BY_SLUG = {t["slug"]: t for t in ALL_TRAININGS}


def grade_training(answers, questions):
    """Grade answers (option indexes or free-text strings) against questions.

    Questions may declare ``type: "text"`` plus an ``accept`` list of valid
    normalized answers; those compare as free text. Everything else compares by
    option index.
    """
    if len(answers) != len(questions):
        raise ValueError(f"expected {len(questions)} answers")
    scored = []
    for i, a in enumerate(answers):
        q = questions[i]
        if q.get("type") == "text":
            accept = {_norm(s) for s in q.get("accept", [])}
            correct = _norm(a) in accept if accept else False
        else:
            try:
                correct = int(a) == q["answer"]
            except (TypeError, ValueError):
                correct = False
        scored.append(correct)
    correct_idx = [i for i, ok in enumerate(scored) if ok]
    missed = sorted({questions[i].get("topic", "General") for i, ok in enumerate(scored) if not ok})
    pct = round(100 * len(correct_idx) / len(questions))
    passed = pct >= 70
    return {
        "total": len(questions),
        "correct": len(correct_idx),
        "percent": pct,
        "passed": passed,
        "missed_topics": missed,
    }


def _norm(s) -> str:
    import unicodedata
    s = unicodedata.normalize("NFKC", str(s or ""))
    s = s.replace("–", "-").replace("—", "-").replace("⁄", "/")
    s = s.replace('"', "").replace("′", "'").replace("“", "").replace("”", "")
    return " ".join(s.strip().lower().split())


def grade_trade_quiz(slug, answers) -> dict:
    """Grade a trade test by slug. Falls back to the OSHA quiz for the 'osha' slug."""
    test = TRAINING_BY_SLUG.get(slug)
    if not test:
        raise ValueError(f"unknown test: {slug}")
    return grade_training(answers, test["questions"])